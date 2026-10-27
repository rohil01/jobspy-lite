"""FastAPI application for JobSpy Lite.

Serves the JSON API and (when built) the React SPA from ``frontend/dist``.
Routes
------
GET  /health                liveness + component status
GET  /                      the SPA (when built)
POST /resume                upload a .docx resume (stored for all runs)
GET  /resume                current resume metadata
GET  /jobs                  list jobs (status filter, auto/manual reject filter, order, pagination)
GET  /jobs/{key}            one job with full details
POST /jobs/{key}/status     accept/reject/reset a job (accept also fills Applications)
POST /jobs/{key}/resume-rewrite   rewrite the stored resume for one job
GET  /jobs/{key}/resume-rewrites  past rewrite attempts for one job
GET  /applications          tracker rows (backfills jobs already accepted)
POST /applications          add a tracker row by hand
POST /applications/import   replace the tracker list
PATCH /applications/{id}    edit a tracker row
DELETE /applications/{id}   delete a tracker row
GET  /stats                 counts by status (incl. manual vs auto rejected) + average score
GET  /runs                  run history (newest first)
POST /run                   start a manual pipeline run (background)
GET  /run/latest            poll the in-flight/last manual run + progress
GET  /scheduler             scheduler status + config
POST /scheduler/config      update cron/interval config (live re-applied)
POST /scheduler/toggle      start/stop the scheduler
GET  /settings              effective settings (config + stored overrides)
POST /settings              persist scrape params / experience window / threshold
GET  /export.xlsx           download the Excel export
"""

import logging
import os
import re
import secrets
import threading
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Any, Dict, Optional

from fastapi import BackgroundTasks, FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from starlette.requests import Request

from . import notify, scheduler, storage
from .agent.resume_io import extract_text, load_resume_docx, save_resume_docx
from .agent.tracing import (
    finish_rewrite,
    job_title,
    rewrite_span,
    setup_tracing,
    shutdown_tracing,
)
from .config import (
    EXCEL_PATH,
    PROJECT_ROOT,
    clamp_int,
    normalize_provider,
    normalize_scrape_counts,
)
from .pipeline import PipelineContext, request_shutdown, run_pipeline, wait_for_idle
from .schemas import (
    ApplicationImport,
    ApplicationPatch,
    ApplicationWrite,
    HealthResponse,
    JobListResponse,
    JobStatusIn,
    PurgeIn,
    ResumeInfo,
    RewriteIn,
    RunLatest,
    SchedulerStatus,
    SettingsIn,
    StatsResponse,
)

MAX_RESUME_BYTES = 2 * 1024 * 1024

logger = logging.getLogger(__name__)


@asynccontextmanager
async def _lifespan(_app: FastAPI):
    """On shutdown, stop taking new scores and wait for the current run."""
    storage.fail_interrupted_rewrites()
    setup_tracing()
    yield
    request_shutdown()
    wait_for_idle(180)
    shutdown_tracing()


app = FastAPI(
    title="JobSpy Lite API",
    version="0.1.0",
    description="Scrape jobs, score them against a stored resume, run on a "
    "schedule, and get Telegram alerts for the best matches.",
    lifespan=_lifespan,
)

# uvicorn app.api:app (Docker) never passes through cli.main. Import-time
# setup covers that path; cli.main calls it again, and both are no-ops once
# a provider exists. The lifespan flushes on shutdown.
setup_tracing()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


def _presented_token(request: Request) -> str:
    header = request.headers.get("x-api-token", "").strip()
    if header:
        return header
    return request.query_params.get("token", "").strip()


@app.middleware("http")
async def require_api_token(request: Request, call_next):
    """Require JOBSPY_API_TOKEN only when that env var is set.

    Unset means the API is open, which is how the Oracle Cloud deploy is
    used: the dashboard is opened from a browser on the public IP, and no
    token is shared. /health, the SPA, and /assets stay open either way so
    the page and the Docker healthcheck can load.
    """
    path = request.url.path
    if (
        request.method == "OPTIONS"
        or path == "/health"
        or path == "/"
        or path.startswith("/assets")
    ):
        return await call_next(request)
    expected = os.environ.get("JOBSPY_API_TOKEN", "").strip()
    if not expected:
        return await call_next(request)
    presented = _presented_token(request)
    if len(presented) == len(expected) and secrets.compare_digest(presented, expected):
        return await call_next(request)
    return _unauthorized("Unauthorized")


def _unauthorized(detail: str) -> JSONResponse:
    response = JSONResponse(status_code=401, content={"detail": detail})
    response.headers["Access-Control-Allow-Origin"] = "*"
    return response

# In-flight manual-run state (process-local; run rows live in SQLite).
_run_state_lock = threading.Lock()
_run_state: Dict[str, Any] = {"running": False, "progress": None, "percent": None,
                              "summary": None, "error": None}


def _resume_hash(text: str) -> str:
    import hashlib

    return hashlib.sha256(" ".join(text.split()).encode("utf-8")).hexdigest()


# --------------------------------------------------------------------------- #
# Health
# --------------------------------------------------------------------------- #
@app.get("/health", response_model=HealthResponse)
def health() -> HealthResponse:
    from .config import AI_MODEL, DB_PATH, TELEGRAM_BOT_TOKEN, TELEGRAM_CHAT_ID

    db_ready = False
    try:
        storage.query("SELECT 1")
        db_ready = True
    except Exception:  # noqa: BLE001 - health must never raise
        db_ready = False
    resume = storage.get_setting("resume") or {}
    return HealthResponse(
        status="ok",
        ai_model=AI_MODEL,
        scheduler_running=scheduler.is_healthy(),
        resume_uploaded=bool(resume.get("text")),
        telegram_configured=bool(TELEGRAM_BOT_TOKEN and TELEGRAM_CHAT_ID),
        db_ready=db_ready,
    )


# --------------------------------------------------------------------------- #
# Resume (step 3)
# --------------------------------------------------------------------------- #
@app.post("/resume", response_model=ResumeInfo)
async def upload_resume(resume: UploadFile = File(...)) -> ResumeInfo:
    """Upload a .docx resume; its text is stored once and used by every run."""
    data = await resume.read()
    if not data:
        raise HTTPException(status_code=400, detail="Empty resume file.")
    if len(data) > MAX_RESUME_BYTES:
        raise HTTPException(status_code=413, detail="Resume must be 2 MB or smaller.")
    try:
        text = extract_text(data)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(
            status_code=400,
            detail=f"Could not read resume. Only valid .docx files are supported ({exc}).",
        )
    if not text.strip():
        raise HTTPException(status_code=400, detail="Resume contained no readable text.")

    fingerprint = _resume_hash(text)
    from datetime import datetime, timezone

    save_resume_docx(data)
    storage.set_setting(
        "resume",
        {"text": text, "name": resume.filename, "hash": fingerprint,
         "uploaded_at": datetime.now(timezone.utc).isoformat()},
    )
    logger.info("Resume stored: name=%s chars=%d", resume.filename, len(text))
    return ResumeInfo(name=resume.filename, chars=len(text), hash=fingerprint)


@app.get("/resume", response_model=ResumeInfo)
def resume_info() -> ResumeInfo:
    resume = storage.get_setting("resume") or {}
    return ResumeInfo(
        name=resume.get("name"),
        chars=len(resume.get("text") or "") or None,
        hash=resume.get("hash"),
        uploaded_at=resume.get("uploaded_at"),
    )


# --------------------------------------------------------------------------- #
# Jobs + accept/reject (step 6)
# --------------------------------------------------------------------------- #
@app.get("/jobs", response_model=JobListResponse)
def list_jobs_endpoint(
    status: Optional[str] = None,
    auto_rejected: Optional[bool] = None,
    ai_failed: bool = False,
    order: str = "score",
    limit: Optional[int] = None,
    offset: int = 0,
    exclude_status: Optional[str] = None,
) -> JobListResponse:
    """List jobs; ``auto_rejected`` narrows rejected rows to the pipeline's
    experience-window rejects (true) or the operator's own rejects (false).

    ``ai_failed`` returns postings that still have no score after a scoring
    failure, including ones past the retry cap.

    ``exclude_status`` is applied before ``limit``, so a page of active jobs
    is not crowded out by rejected ones.
    """
    if status and status not in ("new", "accepted", "rejected"):
        raise HTTPException(status_code=400, detail="status must be new/accepted/rejected.")
    if exclude_status and exclude_status not in ("new", "accepted", "rejected"):
        raise HTTPException(status_code=400, detail="exclude_status must be new/accepted/rejected.")
    jobs = storage.list_jobs(
        status=status, limit=limit, offset=offset, order=order,
        auto_rejected=auto_rejected, ai_failed=ai_failed,
        exclude_status=exclude_status,
    )
    counts = storage.counts_by_status()
    filtered_total = storage.count_jobs(
        status=status, auto_rejected=auto_rejected, ai_failed=ai_failed,
        exclude_status=exclude_status,
    )
    return JobListResponse(
        count=len(jobs),
        total=counts.get("total", 0),
        filtered_total=filtered_total,
        jobs=jobs,
    )


def _rewrite_context(job: Dict[str, Any]) -> Dict[str, str]:
    from .config import EXPERIENCE_MAX_YEARS, EXPERIENCE_MIN_YEARS, scrape_defaults

    min_years = storage.get_setting("experience_min_years", EXPERIENCE_MIN_YEARS)
    max_years = storage.get_setting("experience_max_years", EXPERIENCE_MAX_YEARS)
    if max_years is None:
        years = f"{min_years}+"
    else:
        years = f"{min_years}-{max_years}"
    scrape = {**scrape_defaults(), **(storage.get_setting("scrape_params") or {})}
    terms = scrape.get("search_terms") or []
    if isinstance(terms, str):
        term_text = terms
    else:
        term_text = ", ".join(str(item) for item in terms)
    industry = f"{job.get('title') or ''}. Search terms: {term_text}".strip()
    return {"years": years, "industry": industry}


def _stamp_rewrite(result: Dict[str, Any], applied_rows: list) -> list:
    """Changes are the whole winning lineage, not only the last try.

    A later try can score higher without new wording. Those earlier patches
    are still what changed the résumé, so they stay in the list the panel shows.
    """
    chain = result.get("patch_chain") or []
    flat = [patch for group in chain for patch in (group or [])]
    flags = [bool(row.get("applied")) for row in applied_rows]
    winner = result.get("winner_try")
    for row in result.get("trace") or []:
        row["winner"] = winner is not None and row.get("try") == winner
    winning: list = []
    for index, patch in enumerate(flat):
        applied = flags[index] if index < len(flags) else False
        winning.append({**patch, "applied": applied})
    return winning


_rewrite_guard = threading.Lock()
_rewrite_locks: Dict[str, threading.Lock] = {}


def _rewrite_lock(job_key: str) -> threading.Lock:
    with _rewrite_guard:
        lock = _rewrite_locks.get(job_key)
        if lock is None:
            lock = threading.Lock()
            _rewrite_locks[job_key] = lock
        return lock


def _rewrite_docx_name(company: str) -> str:
    """Rohil_Chaturvedi_Infosys.docx from the job's company."""
    cleaned = re.sub(r"[^A-Za-z0-9]+", "_", (company or "").strip()).strip("_")
    return f"Rohil_Chaturvedi_{cleaned or 'Company'}.docx"


def _rewrite_dir() -> Path:
    from .config import DB_PATH

    folder = DB_PATH.parent / "rewrites"
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def _resolve_rewrite_path(stored: Optional[str]) -> Optional[Path]:
    """Find a rewrite file by absolute path or by basename under data/rewrites."""
    if not stored:
        return None
    path = Path(stored)
    if path.is_file():
        return path
    local = _rewrite_dir() / path.name
    if local.is_file():
        return local
    return None


def _save_rewrite_docx(
    job: Optional[Dict[str, Any]],
    original: bytes,
    patch_chain: Optional[list] = None,
) -> tuple:
    """Write the patched résumé, or a copy of the original when nothing applied."""
    from .agent.resume_io import enforce_one_page

    applied_rows: list = []
    patched = original
    if patch_chain:
        patched, applied_rows = enforce_one_page(original, patch_chain)
    path = _rewrite_dir() / _rewrite_docx_name((job or {}).get("company") or "")
    path.write_bytes(patched)
    return str(path), applied_rows


def _flatten_patches(patch_chain: Optional[list] = None, winning_patches: Optional[list] = None) -> list:
    patches: list = []
    if patch_chain:
        for group in patch_chain:
            for patch in group or []:
                if patch:
                    patches.append(patch)
    elif winning_patches:
        patches = [patch for patch in winning_patches if patch]
    return patches


def _apply_text_patches(text: str, patches: Optional[list] = None) -> str:
    current = text or ""
    for patch in patches or []:
        original = str((patch or {}).get("original") or "")
        replacement = str((patch or {}).get("replacement") or "")
        if original and original in current:
            current = current.replace(original, replacement, 1)
    return current


def _docx_from_text(text: str) -> bytes:
    from io import BytesIO

    from docx import Document

    document = Document()
    lines = (text or "").split("\n")
    if not lines:
        document.add_paragraph("")
    else:
        for line in lines:
            document.add_paragraph(line)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _write_available_docx(
    job: Optional[Dict[str, Any]],
    patch_chain: Optional[list] = None,
    winning_patches: Optional[list] = None,
) -> tuple:
    """Prefer the uploaded Word file; otherwise build one from stored resume text."""
    original = load_resume_docx()
    if original:
        return _save_rewrite_docx(job, original, patch_chain)
    resume = storage.get_setting("resume") or {}
    text = resume.get("text") or ""
    if not text.strip():
        return None, []
    patched = _apply_text_patches(text, _flatten_patches(patch_chain, winning_patches))
    path = _rewrite_dir() / _rewrite_docx_name((job or {}).get("company") or "")
    path.write_bytes(_docx_from_text(patched))
    return str(path), []


def _ensure_rewrite_docx(row: Optional[Dict[str, Any]]) -> Optional[Dict[str, Any]]:
    """Fill in a finished rewrite that never got a file."""
    if row is None or row.get("status") == "running":
        return row
    if _resolve_rewrite_path(row.get("docx_path")):
        return row
    job = storage.get_job(row.get("job_key") or "") or {}
    path, _applied = _write_available_docx(
        job, row.get("patch_chain"), row.get("winning_patches")
    )
    if not path:
        return row
    return storage.update_resume_rewrite(row["id"], {"docx_path": path}) or {
        **row,
        "docx_path": path,
        "docx_ready": True,
    }


def _finish_rewrite(rewrite_id: str, job_key: str, payload: RewriteIn) -> None:
    """Run the loop and fill the row that was inserted as ``running``."""
    from .agent.resume_rewrite import make_client, run_loop
    from .config import agent_defaults

    try:
        job = storage.get_job(job_key)
        resume = storage.get_setting("resume") or {}
        resume_text = resume.get("text") or ""
        agent = {**agent_defaults(), **(storage.get_setting("agent_params") or {})}
        max_tries = clamp_int(
            payload.max_tries if payload.max_tries is not None else agent.get("rewrite_max_tries"),
            1, 5, 3,
        )
        additional = (
            storage.get_setting("additional_info") or ""
            if payload.additional_info is None
            else payload.additional_info
        )
        context = _rewrite_context(job or {})
        client = make_client(payload.model or "", payload.base_url or "", agent)

        def complete(messages, text_format, label):
            return client.parse(messages, text_format, label=label)

        original = load_resume_docx()
        title = job_title(job)
        trace_inputs = {
            "job": job,
            "resume_text": resume_text,
            "additional_info": additional or "",
            "years": context["years"],
            "industry": context["industry"],
            "company_type": payload.company_type or "",
            "max_tries": max_tries,
        }
        with rewrite_span(
            rewrite_id, job_key, max_tries, title=title, inputs=trace_inputs
        ) as trace_span:
            result = run_loop(
                resume_text=resume_text,
                job=job or {},
                additional_info=additional or "",
                max_tries=max_tries,
                complete=complete,
                years=context["years"],
                industry=context["industry"],
                company_type=payload.company_type or "",
                resume_docx=original,
            )
            docx_path, applied_rows = _write_available_docx(
                job, result.get("patch_chain")
            )
            winning = _stamp_rewrite(result, applied_rows or [])
            failed = bool(result.get("error"))
            status = "error" if failed else "done"
            storage.update_resume_rewrite(
                rewrite_id,
                {
                    "trace": result["trace"],
                    "winning_total": result["winning_total"],
                    "winning_patches": winning,
                    "patch_chain": result["patch_chain"],
                    "docx_path": docx_path,
                    "error": result["error"],
                    "status": status,
                },
            )
            finish_rewrite(
                trace_span,
                status,
                output={
                    "status": status,
                    "winning_total": result["winning_total"],
                    "winning_patches": winning,
                    "docx_path": docx_path,
                    "error": result["error"],
                },
            )
    except Exception as exc:  # noqa: BLE001 - the row must leave the running state
        logger.exception("Resume rewrite failed for %s", job_key)
        patch = {"status": "error", "error": str(exc)}
        failed_job = storage.get_job(job_key) or {}
        path, _applied = _write_available_docx(failed_job)
        if path:
            patch["docx_path"] = path
        storage.update_resume_rewrite(rewrite_id, patch)


@app.post("/jobs/{job_key:path}/resume-rewrite", status_code=202)
def rewrite_resume(job_key: str, payload: RewriteIn) -> Dict[str, Any]:
    """Start a resume rewrite. The row is filled in when the loop finishes."""
    import uuid

    from .agent.resume_rewrite import rewrite_target
    from .config import agent_defaults

    job = storage.get_job(job_key)
    if job is None:
        raise HTTPException(status_code=404, detail=f"Unknown job '{job_key}'.")
    if not (job.get("description") or "").strip():
        raise HTTPException(
            status_code=400,
            detail="This job has no description to rewrite against.",
        )
    resume = storage.get_setting("resume") or {}
    if not (resume.get("text") or "").strip():
        raise HTTPException(status_code=400, detail="Upload a resume before rewriting.")

    agent = {**agent_defaults(), **(storage.get_setting("agent_params") or {})}
    max_tries = clamp_int(
        payload.max_tries if payload.max_tries is not None else agent.get("rewrite_max_tries"),
        1, 5, 3,
    )
    additional = (
        storage.get_setting("additional_info") or ""
        if payload.additional_info is None
        else payload.additional_info
    )
    target = rewrite_target(payload.model or "", payload.base_url or "", agent)
    with _rewrite_lock(job_key):
        current = storage.list_resume_rewrites(job_key)
        if current and current[0].get("status") == "running":
            return current[0]
        pending = storage.insert_resume_rewrite(
            {
                "id": str(uuid.uuid4()),
                "job_key": job_key,
                "model": target["model"],
                "base_url": target["base_url"],
                "resume_hash": resume.get("hash"),
                "max_tries": max_tries,
                "trace": [],
                "winning_total": None,
                "winning_patches": [],
                "patch_chain": [],
                "docx_path": None,
                "error": None,
                "additional_info": additional or "",
                "company_type": payload.company_type or "",
                "status": "running",
            }
        )
    threading.Thread(
        target=_finish_rewrite,
        args=(pending["id"], job_key, payload),
        daemon=True,
    ).start()
    return pending


@app.get("/jobs/{job_key:path}/resume-rewrites")
def list_resume_rewrites_endpoint(job_key: str) -> Dict[str, Any]:
    if storage.get_job(job_key) is None:
        raise HTTPException(status_code=404, detail=f"Unknown job '{job_key}'.")
    return {
        "rewrites": [
            _ensure_rewrite_docx(row) for row in storage.list_resume_rewrites(job_key)
        ]
    }


@app.get("/jobs/{job_key:path}/resume-rewrites/{rewrite_id}/docx")
def download_rewritten_resume(job_key: str, rewrite_id: str) -> FileResponse:
    row = storage.get_resume_rewrite(rewrite_id)
    if row is None or row.get("job_key") != job_key:
        raise HTTPException(status_code=404, detail="No patched resume for this attempt.")
    row = _ensure_rewrite_docx(row)
    path = _resolve_rewrite_path((row or {}).get("docx_path"))
    if path is None:
        raise HTTPException(status_code=404, detail="No patched resume for this attempt.")
    job = storage.get_job(job_key) or {}
    return FileResponse(
        path,
        filename=_rewrite_docx_name(job.get("company") or ""),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


@app.get("/jobs/{job_key:path}")
def get_job_endpoint(job_key: str) -> Dict[str, Any]:
    job = storage.get_job(job_key)
    if job is None:
        raise HTTPException(status_code=404, detail=f"Unknown job '{job_key}'.")
    return job


@app.post("/jobs/{job_key:path}/status")
def set_job_status_endpoint(job_key: str, payload: JobStatusIn) -> Dict[str, Any]:
    try:
        updated = storage.set_job_status(job_key, payload.new_status)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    if not updated:
        raise HTTPException(status_code=404, detail=f"Unknown job '{job_key}'.")
    if payload.new_status == "accepted":
        job = storage.get_job(job_key)
        if job is not None:
            storage.ensure_application_for_job(job)
    return {"job_key": job_key, "status": payload.new_status}


@app.get("/applications")
def list_applications() -> Dict[str, Any]:
    return {"applications": storage.list_applications()}


@app.post("/applications", status_code=201)
def create_application(payload: ApplicationWrite) -> Dict[str, Any]:
    try:
        return storage.create_application(payload.model_dump())
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))


@app.post("/applications/import")
def import_applications(payload: ApplicationImport) -> Dict[str, Any]:
    try:
        rows = storage.replace_applications(
            [item.model_dump() for item in payload.applications]
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    return {"applications": rows}


@app.patch("/applications/{app_id}")
def update_application(app_id: str, payload: ApplicationPatch) -> Dict[str, Any]:
    try:
        updated = storage.update_application(
            app_id, payload.model_dump(exclude_unset=True)
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    if updated is None:
        raise HTTPException(status_code=404, detail=f"Unknown application '{app_id}'.")
    return updated


@app.delete("/applications/{app_id}")
def delete_application(app_id: str) -> Dict[str, Any]:
    if not storage.delete_application(app_id):
        raise HTTPException(status_code=404, detail=f"Unknown application '{app_id}'.")
    return {"id": app_id, "deleted": True}


@app.get("/stats", response_model=StatsResponse)
def stats() -> StatsResponse:
    counts = storage.counts_by_status()
    return StatsResponse(
        total=counts.get("total", 0),
        new=counts.get("new", 0),
        accepted=counts.get("accepted", 0),
        rejected=counts.get("rejected", 0),
        manual_rejected=counts.get("manual_rejected", 0),
        auto_rejected=counts.get("auto_rejected", 0),
        avg_score=counts.get("avg_score"),
    )


# --------------------------------------------------------------------------- #
# Runs + manual trigger (steps 1/4)
# --------------------------------------------------------------------------- #
@app.post("/run", status_code=202)
def start_run(
    background: BackgroundTasks, force_rescore: bool = False
) -> Dict[str, Any]:
    """Start a manual pipeline run in the background; poll /run/latest."""
    with _run_state_lock:
        if _run_state.get("running"):
            raise HTTPException(status_code=409, detail="A run is already in progress.")
        _run_state.update(running=True, progress="queued", percent=0,
                          summary=None, error=None)

    def _task() -> None:
        def _progress(message: str, percent: Optional[float]) -> None:
            with _run_state_lock:
                _run_state["progress"] = message
                _run_state["percent"] = percent

        try:
            summary = run_pipeline(
                trigger="manual", force_rescore=force_rescore,
                progress_callback=_progress,
            )
            with _run_state_lock:
                _run_state.update(running=False, summary=summary,
                                  error=summary.get("error"))
        except Exception as exc:  # noqa: BLE001 - surfaced via /run/latest
            logger.exception("Manual run crashed")
            with _run_state_lock:
                _run_state.update(running=False, error=str(exc))

    background.add_task(_task)
    return {"started": True}


@app.get("/run/latest", response_model=RunLatest)
def run_latest() -> RunLatest:
    with _run_state_lock:
        state = dict(_run_state)
    last_db = storage.list_runs(limit=1)
    return RunLatest(
        running=state.get("running", False),
        progress=state.get("progress"),
        percent=state.get("percent"),
        summary=state.get("summary"),
        error=state.get("error"),
        last_run=last_db[0] if last_db else None,
    )


@app.get("/runs")
def runs(limit: int = 50) -> Dict[str, Any]:
    return {"runs": storage.list_runs(limit=limit)}


# --------------------------------------------------------------------------- #
# Scheduler (step 4)
# --------------------------------------------------------------------------- #
@app.get("/scheduler", response_model=SchedulerStatus)
def scheduler_status() -> SchedulerStatus:
    return SchedulerStatus(**scheduler.status())


@app.post("/scheduler/config", response_model=SchedulerStatus)
def scheduler_config(config: Dict[str, Any]) -> SchedulerStatus:
    try:
        scheduler.set_config(config)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    return SchedulerStatus(**scheduler.status())


@app.post("/scheduler/toggle", response_model=SchedulerStatus)
def scheduler_toggle(enabled: bool) -> SchedulerStatus:
    scheduler.toggle(enabled)
    return SchedulerStatus(**scheduler.status())


# --------------------------------------------------------------------------- #
# Settings
# --------------------------------------------------------------------------- #
@app.get("/settings")
def get_settings() -> Dict[str, Any]:
    from .config import AI_MODEL, REWRITE_MODEL, agent_defaults, scrape_defaults

    context = PipelineContext()
    return {
        "scrape_params": normalize_scrape_counts(
            {**scrape_defaults(), **(storage.get_setting("scrape_params") or {})}
        ),
        "agent_params": {**agent_defaults(), **(storage.get_setting("agent_params") or {})},
        "rewrite_model_default": REWRITE_MODEL or AI_MODEL,
        "experience_min_years": context.experience_min,
        "experience_max_years": context.experience_max,
        "score_threshold": context.score_threshold,
        "max_workers": context.max_workers,
        "retention_days": context.retention_days,
        "additional_info": storage.get_setting("additional_info") or "",
    }


@app.post("/settings")
def update_settings(settings: SettingsIn) -> Dict[str, Any]:
    if settings.scrape_params is not None:
        params = normalize_scrape_counts(dict(settings.scrape_params))
        storage.set_setting("scrape_params", params)
    if settings.agent_params is not None:
        params = dict(settings.agent_params)
        # API keys are env/server-only, never dashboard-editable.
        params.pop("api_key", None)
        model = str(params.get("ai_model") or "").strip()
        if not model:
            params.pop("ai_model", None)
        base_url = str(params.get("ai_base_url") or "").strip()
        params["ai_base_url"] = base_url  # empty string = fall back to env default
        raw_provider = str(params.get("ai_provider") or "").strip().lower()
        if raw_provider == "custom":
            params["ai_provider"] = "custom"
        else:
            provider = normalize_provider(raw_provider)
            if provider:
                params["ai_provider"] = provider
            else:
                params.pop("ai_provider", None)
        rate = params.get("ai_rate_limit_per_min")
        try:
            params["ai_rate_limit_per_min"] = max(0, int(rate)) if rate is not None else 30
        except (TypeError, ValueError):
            params["ai_rate_limit_per_min"] = 30
        if "max_workers" in params:
            params["max_workers"] = clamp_int(params.get("max_workers"), 1, 8, 4)
        params["rewrite_model"] = str(params.get("rewrite_model") or "").strip()
        params["rewrite_base_url"] = str(params.get("rewrite_base_url") or "").strip()
        if "rewrite_max_tries" in params:
            params["rewrite_max_tries"] = clamp_int(params.get("rewrite_max_tries"), 1, 5, 3)
        storage.set_setting("agent_params", params)
    if settings.experience_min_years is not None:
        storage.set_setting("experience_min_years", settings.experience_min_years)
    if settings.experience_max_years is not None:
        storage.set_setting("experience_max_years", settings.experience_max_years)
    if settings.score_threshold is not None:
        storage.set_setting("score_threshold", settings.score_threshold)
    if settings.retention_days is not None:
        # 0 = keep everything; negative values are clamped to 0.
        storage.set_setting("retention_days", max(0, int(settings.retention_days)))
    if "additional_info" in settings.model_fields_set:
        storage.set_setting("additional_info", settings.additional_info or "")
    return get_settings()


# --------------------------------------------------------------------------- #
# Excel export (step 1)
# --------------------------------------------------------------------------- #
@app.get("/export.xlsx")
def export_xlsx() -> FileResponse:
    from .excel import export_excel

    path = export_excel()
    if not path.exists():
        raise HTTPException(status_code=500, detail="Excel export failed.")
    return FileResponse(path, filename="jobs.xlsx",
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# --------------------------------------------------------------------------- #
# Notification test
# --------------------------------------------------------------------------- #
@app.post("/maintenance/purge")
def purge_old_data(payload: PurgeIn) -> Dict[str, int]:
    """Delete jobs and run history older than N days. Accepted jobs stay."""
    if payload.older_than_days < 1:
        raise HTTPException(status_code=400, detail="older_than_days must be >= 1.")
    try:
        return storage.purge_older_than(payload.older_than_days)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))


@app.post("/notify/test")
def notify_test() -> Dict[str, Any]:
    """Send any pending above-threshold matches to Telegram right now."""
    context = PipelineContext()
    sent = notify.notify_above_threshold(context.score_threshold)
    return {"sent": sent}


# --------------------------------------------------------------------------- #
# Static SPA (built frontend)
# --------------------------------------------------------------------------- #
_DIST = PROJECT_ROOT / "frontend" / "dist"
if (_DIST / "index.html").exists():  # pragma: no cover - build-dependent
    app.mount("/assets", StaticFiles(directory=_DIST / "assets"), name="assets")

    @app.get("/", include_in_schema=False)
    def index() -> FileResponse:
        return FileResponse(_DIST / "index.html")
else:

    @app.get("/", include_in_schema=False)
    def index_dev() -> JSONResponse:
        return JSONResponse(
            {
                "message": "JobSpy Lite API is running. The UI is not built yet — "
                "run `npm run build` in frontend/ or use the Vite dev server."
            }
        )
