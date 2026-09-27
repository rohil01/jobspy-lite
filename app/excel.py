"""Excel export: the jobs table rendered into a formatted ``.xlsx`` workbook."""

import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import EXCEL_PATH

logger = logging.getLogger(__name__)

_COLUMNS = [
    ("Job Key", 46),
    ("Title", 38),
    ("Company", 24),
    ("Location", 22),
    ("Site", 12),
    ("Score", 8),
    ("Verdict", 10),
    ("Experience Match", 16),
    ("Required Years", 16),
    ("Matched Skills", 40),
    ("Missing Skills", 40),
    ("Status", 10),
    ("Notified", 9),
    ("Reasoning", 50),
    ("Date Posted", 14),
    ("URL", 50),
    ("First Seen", 24),
    ("Last Seen", 24),
]

_HEADER_FILL = PatternFill("solid", fgColor="1F2933")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")


def _required_years_text(job: Dict[str, Any]) -> str:
    low, high = job.get("required_years_min"), job.get("required_years_max")
    if low is None and high is None:
        return ""
    if high is None:
        return f"{low}+"
    if low == high:
        return str(low)
    return f"{low}-{high}"


def write_excel(jobs: List[Dict[str, Any]], output_path: Any = None) -> Path:
    """Write ``jobs`` (storage-style dicts) to an .xlsx file and return its path.

    ``output_path`` defaults to ``data/jobs.xlsx`` from config; tests pass a
    temporary path.
    """
    path = Path(output_path) if output_path else EXCEL_PATH
    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Jobs"

    for index, (header, width) in enumerate(_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _WRAP
        sheet.column_dimensions[get_column_letter(index)].width = width

    for row_number, job in enumerate(jobs, start=2):
        matched = ", ".join(job.get("matched_skills") or [])
        missing = ", ".join(job.get("missing_skills") or [])
        row = [
            job.get("job_key"),
            job.get("title"),
            job.get("company"),
            job.get("location"),
            job.get("site"),
            job.get("score"),
            job.get("verdict"),
            "yes" if job.get("experience_match") else "no",
            _required_years_text(job),
            matched,
            missing,
            job.get("status"),
            "yes" if job.get("notified") else "no",
            job.get("reasoning"),
            job.get("date_posted"),
            job.get("job_url"),
            job.get("first_seen_at"),
            job.get("last_seen_at"),
        ]
        for column_number, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_number, column=column_number, value=value)
            cell.alignment = _WRAP

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(path)
    logger.info("Wrote %d jobs to %s", len(jobs), path)
    return path


def export_excel(output_path: Any = None) -> Path:
    """Export the full jobs table to Excel (called after runs and on demand)."""
    from . import storage

    jobs = storage.list_jobs(order="score")
    return write_excel(jobs, output_path)
