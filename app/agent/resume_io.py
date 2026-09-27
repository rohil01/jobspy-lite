"""Resume document I/O (``.docx`` only, via python-docx).

``extract_text`` reads an uploaded resume into plain text for the agents.
Legacy binary ``.doc`` files are not supported — convert them to ``.docx`` first.
"""

from io import BytesIO
from typing import List

from docx import Document


def extract_text(data: bytes) -> str:
    """Extract plain text from ``.docx`` bytes (paragraphs and table cells)."""
    document = Document(BytesIO(data))

    lines: List[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)

    # Include table content (many resumes use tables for layout).
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    return "\n".join(lines)
