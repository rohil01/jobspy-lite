"""Excel export tests."""

from app.excel import write_excel


def test_write_excel_produces_valid_xlsx(tmp_path):
    jobs = [
        {
            "job_key": "https://x/1",
            "title": "AI Engineer",
            "company": "Acme",
            "location": "Bengaluru",
            "site": "linkedin",
            "score": 85,
            "verdict": "strong",
            "experience_match": 1,
            "required_years_min": 1,
            "required_years_max": 2,
            "matched_skills": ["python", "sql"],
            "missing_skills": ["go"],
            "status": "accepted",
            "notified": 1,
            "reasoning": "good fit",
            "job_url": "https://x/1",
            "first_seen_at": "2026-09-26T00:00:00+00:00",
            "last_seen_at": "2026-09-26T00:00:00+00:00",
        },
        {
            "job_key": "https://x/2",
            "title": "B",
            "score": None,
        },
    ]
    path = write_excel(jobs, tmp_path / "out.xlsx")
    assert path.exists()
    assert path.read_bytes()[:2] == b"PK"

    from openpyxl import load_workbook

    workbook = load_workbook(path)
    sheet = workbook.active
    assert sheet.cell(row=1, column=2).value == "Title"
    assert sheet.cell(row=2, column=2).value == "AI Engineer"
    assert sheet.cell(row=2, column=10).value == "python, sql"
    assert sheet.cell(row=3, column=2).value == "B"
    assert sheet.max_row == 3
